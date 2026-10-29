# -*- coding: utf-8 -*-
"""Reopen yuzpalang-showcase.pptx and verify the deck contract structurally."""
from pptx import Presentation
from pptx.util import Emu
from pptx.oxml.ns import qn

prs = Presentation("yuzpalang-showcase.pptx")
print("slide count:", len(prs.slides._sldIdLst))
print("size (in):", round(prs.slide_width / 914400, 3), "x", round(prs.slide_height / 914400, 3))
print("core title:", prs.core_properties.title)
print("core author:", prs.core_properties.author)

issues = []
fa_count = 0
rtl_count = 0
pic_alt = []
notes_count = 0
has_chart = False
has_table = False

for idx, slide in enumerate(prs.slides, start=1):
    for shape in slide.shapes:
        # images + alt text
        if shape.shape_type == 13:  # PICTURE
            nv = shape._element.nvPicPr
            name = nv.cNvPr.get("name")
            descr = nv.cNvPr.get("descr")
            pic_alt.append((name, descr))
        if shape.has_chart:
            has_chart = True
        if shape.has_table:
            has_table = True
            # Robust RTL table layout is achieved by explicit column ordering:
            # trait/rank labels live in the RIGHT-most column (col index 1).
            t = shape.table
            header_vals = [t.cell(0, j).text for j in range(len(t.columns))]
            if header_vals != ["ارزش", "ویژگی"]:
                issues.append(f"slide {idx}: table column order unexpected: {header_vals}")
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                pPr = p._p.find(qn("a:pPr"))
                if pPr is not None and pPr.get("rtl") == "1":
                    rtl_count += 1
                for r in p.runs:
                    rPr = r._r.find(qn("a:rPr"))
                    if rPr is not None:
                        lang = rPr.get("lang")
                        if lang and lang.upper().startswith("FA"):
                            fa_count += 1
                        cs = rPr.find(qn("a:cs"))
                        if cs is not None and cs.get("typeface") != "Vazirmatn":
                            issues.append(f"slide {idx}: cs font not Vazirmatn")
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
        notes_count += 1

print("FARSI-tagged runs:", fa_count)
print("rtl=\"1\" paragraphs:", rtl_count)
print("slides with speaker notes:", notes_count)
print("chart present:", has_chart)
print("table present:", has_table)
print("images (name / alt):")
for n, d in pic_alt:
    print("   -", n, "/", d)
    if not d:
        issues.append("missing alt text: " + n)

# check a few paragraphs actually contain Persian digits
import re
total_persian = 0
for slide in prs.slides:
    for shape in slide.shapes:
        if shape.has_text_frame:
            txt = shape.text_frame.text
            total_persian += len(re.findall(r"[\u06F0-\u06F9]", txt))
print("Persian digits found (approx count of chars):", total_persian)

print("\nISSUES:" if issues else "\nNO ISSUES")
for i in issues:
    print("  !", i)
