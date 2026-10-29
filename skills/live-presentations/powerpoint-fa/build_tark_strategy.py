# -*- coding: utf-8 -*-
"""Build a Persian RTL deck on an underrated strategy topic: ترک استراتژیک.

Topic: Planned Abandonment — what to STOP (Drucker). Chosen because Persian
strategy discourse over-covers vision/growth/OKR and under-covers stopping.
All numbers are framework/educational (labeled آموزشی); no fake citations.

Canon: references/slide-rules-fa.md. Patterns: archetype-catalog-fa.md.
Theme + RTL primitives: scripts/fa_helpers.py (single source).
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent / "scripts"))

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from fa_helpers import (
    SLIDE_W, SLIDE_H, MARGIN, FOOTER_Y, SAND, INK, SECONDARY, ACCENT, TEAL,
    CARD, CARDBORDER, FAINT, DARKPANEL, FONT, set_run_font, set_para_rtl,
    set_bullet,
)

OUT = "tark-strategy-fa.pptx"
TOTAL_FA = "۷"

FA_DIGITS = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")


def fa(n):
    return str(n).translate(FA_DIGITS)


prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
blank = prs.slide_layouts[6]

cp = prs.core_properties
cp.title = "ترک استراتژیک — رشد با توقف آغاز می‌شود"
cp.subject = "ارائه آموزشی فارسی و راست‌به‌چپ درباره ترک برنامه‌ریزی‌شده"
cp.author = "powerpoint-fa skill"
cp.keywords = "pptx, persian, farsi, rtl, strategy, ترک استراتژیک"


# ----------------------------------------------------------------------------
# Builders (rect default §5.1; filled shapes carry no border §5.3)
# ----------------------------------------------------------------------------
def add_rect(slide, x, y, w, h, fill):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tb, tf


def para(tf, runs, size=18, bold=False, color=INK, first=False, rtl=True,
         align=None, bullet=None, space_after=6, space_before=0, line=None):
    if isinstance(runs, str):
        runs = [(runs, {})]
    p = tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.alignment = align if align is not None else (PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT)
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    if line is not None:
        p.line_spacing = line
    set_para_rtl(p, rtl)
    if bullet is True:
        set_bullet(p, level=0, char="•")
    elif bullet is False:
        set_bullet(p, char=None)
    for text, ov in runs:
        r = p.add_run()
        r.text = text
        set_run_font(
            r,
            size=ov.get("size", size),
            bold=ov.get("bold", bold),
            color=ov.get("color", color),
            font=ov.get("font", FONT),
        )
    return p


def bg(slide, color=SAND):
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, color)


def header(slide, title, kicker, num=None):
    add_rect(slide, 0, 0, SLIDE_W, Inches(0.14), ACCENT)
    _, tf = textbox(slide, MARGIN, Inches(0.42), SLIDE_W - 2 * MARGIN, Inches(1.1))
    para(tf, kicker, size=13, bold=True, color=TEAL, first=True, space_after=2)
    para(tf, title, size=28, bold=True, color=INK, space_after=0)
    add_rect(slide, MARGIN, Inches(7.06), SLIDE_W - 2 * MARGIN, Pt(1), CARDBORDER)
    # footer: narrow brand box at left edge + narrow number box at right edge
    _, ftf = textbox(slide, MARGIN, Inches(7.12), Inches(3.0), Inches(0.3))
    para(ftf, "ترک استراتژیک", size=10, color=FAINT, align=PP_ALIGN.LEFT, first=True)
    if num is not None:
        _, ntf = textbox(slide, SLIDE_W - MARGIN - Inches(1.1), Inches(7.12),
                         Inches(1.1), Inches(0.3))
        para(ntf, f"{fa(num)} / {TOTAL_FA}", size=10, color=FAINT,
             align=PP_ALIGN.RIGHT, first=True)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def right_strip_card(slide, x, y, w, h, title, body, strip=TEAL,
                     title_size=20, body_size=14):
    """White card (no border §5.3) + accent strip on the RIGHT (RTL start).

    The strip is drawn BEFORE the textbox so color sits behind text.
    """
    add_rect(slide, x, y, w, h, CARD)
    sw = Inches(0.16)
    add_rect(slide, x + w - sw, y + Inches(0.12), sw, h - Inches(0.24), strip)
    _, tf = textbox(slide, x + Inches(0.35), y + Inches(0.22),
                    w - Inches(0.7), h - Inches(0.44))
    para(tf, title, size=title_size, bold=True, color=INK, first=True,
         space_after=6, bullet=False)
    para(tf, body, size=body_size, color=SECONDARY, line=1.15, space_after=0,
         bullet=False)
    return tf


def set_chart_font(chart, font=FONT):
    cs = chart._chartSpace
    for tag in ("a:latin", "a:ea", "a:cs"):
        for el in cs.iter(qn(tag)):
            el.set("typeface", font)
    for tag in ("a:defRPr", "a:endParaRPr"):
        for rPr in cs.iter(qn(tag)):
            for f in ("a:latin", "a:ea", "a:cs"):
                child = rPr.find(qn(f))
                if child is None:
                    child = rPr.makeelement(qn(f), {})
                    rPr.append(child)
                child.set("typeface", font)


# ----------------------------------------------------------------------------
# S1 — Cover (unnumbered)
# ----------------------------------------------------------------------------
s = prs.slides.add_slide(blank)
bg(s)
add_rect(s, 0, 0, SLIDE_W, Inches(0.14), ACCENT)
add_rect(s, MARGIN, Inches(1.5), SLIDE_W - 2 * MARGIN, Inches(4.4), DARKPANEL)
_, tf = textbox(s, MARGIN + Inches(0.7), Inches(1.9),
                SLIDE_W - 2 * MARGIN - Inches(1.4), Inches(3.8))
para(tf, "ترک استراتژیک · موضوع مغفول برنامه‌ریزی", size=15, bold=True,
     color=ACCENT, first=True, space_after=12)
para(tf, "رشد با توقف آغاز می‌شود", size=44, bold=True, color=SAND,
     space_after=8)
para(tf, "کار کمتر بر عهده بگیرید تا اثر بیشتری بگذارید", size=20, bold=True,
     color=CARD, space_after=16)
para(tf, "شماری از سازمان‌ها به افزودن برنامه‌های تازه عادت دارند و به توقف فکر نمی‌کنند. "
         "این ارائه فرایند ترک برنامه‌ریزی‌شده را نشان می‌دهد.",
     size=15, color=CARD, line=1.3, space_after=14)
para(tf, "بر پایه مفهوم «ترک برنامه‌ریزی‌شده» از پیتر دراکر", size=12, color=FAINT)
notes(s, "اسلاید جلد. پیام: رشد با توقف آغاز می‌شود. مخاطب مدیران ارشد و میانی. "
         "مدت ارائه بیست دقیقه. بدون شماره صفحه.")

# ----------------------------------------------------------------------------
# S2/p1 — Three locks (stat cards, uniform brand strips)
# ----------------------------------------------------------------------------
s = prs.slides.add_slide(blank)
bg(s)
header(s, "کار کم‌بازده منابع را قفل می‌کند", "ترک استراتژیک · قفل منابع", num=1)
locks = [
    ("بودجه تکراری", "بودجه سال بعد برابر سال پیش بسته می‌شود"),
    ("زمان قفل‌شده", "ساعت‌ها صرف جلسه‌ها و گزارش‌های کم‌اثر می‌شوند"),
    ("توجه پراکنده", "تمرکز مدیران میان ده‌ها کار پخش می‌شود"),
]
cw = (SLIDE_W - 2 * MARGIN - 2 * Inches(0.3)) / 3
cy, ch = Inches(1.95), Inches(4.5)
for i, (t, b) in enumerate(locks):
    # RTL: first item at highest x (rightmost)
    x = SLIDE_W - MARGIN - (i + 1) * cw - i * Inches(0.3)
    right_strip_card(s, x, cy, cw, ch, t, b, body_size=15)
_, stf = textbox(s, MARGIN, Inches(6.72), Inches(5.0), Inches(0.28))
para(stf, "قاب مفهومی آموزشی", size=11, color=FAINT, first=True)
notes(s, "سه قفل منابع: بودجه و زمان و توجه. این قاب آموزشی است و آمار واقعی نیست.")

# ----------------------------------------------------------------------------
# S3/p2 — Axis table (rows = دسته, cols = دیدگاه; RTL-leading col rightmost)
# ----------------------------------------------------------------------------
s = prs.slides.add_slide(blank)
bg(s)
header(s, "نه هر فعالیتی ماندنی است", "جدول محور · نشانه و اقدام", num=2)
tx, ty, tw, th = MARGIN, Inches(1.85), SLIDE_W - 2 * MARGIN, Inches(4.8)
shape = s.shapes.add_table(4, 3, tx, ty, tw, th)
table = shape.table
table.horz_banding = False
table.columns[0].width = Inches(4.0)    # LEFT: اقدام نخست
table.columns[1].width = Inches(4.9)    # MIDDLE: نشانه
table.columns[2].width = Inches(3.23)   # RIGHT: دسته (row axis, RTL-leading)
table.rows[0].height = Inches(1.0)
for r in range(1, 4):
    table.rows[r].height = Inches(1.27)


def cell(r_, c_, text, size, bold, color, fill):
    cl = table.cell(r_, c_)
    cl.fill.solid()
    cl.fill.fore_color.rgb = fill
    cl.margin_left = Inches(0.15)
    cl.margin_right = Inches(0.15)
    cl.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_ = cl.text_frame
    tf_.word_wrap = True
    p_ = tf_.paragraphs[0]
    p_.alignment = PP_ALIGN.RIGHT
    set_para_rtl(p_, True)
    set_bullet(p_, char=None)
    run_ = p_.add_run()
    run_.text = text
    set_run_font(run_, size=size, bold=bold, color=color)


headers = ["اقدام نخست", "نشانه", "دسته"]
for j, h in enumerate(headers):
    cell(0, j, h, 19, True, INK, CARD)
rows = [
    ("تقویت و تمرکز", "تقاضای پایدار و مزیت روشن دارد", "نگهداشت"),
    ("کوچک‌سازی و بازتعریف", "منابع زیادی مصرف می‌کند ولی اثر متوسط دارد", "بازطراحی"),
    ("توقف و بازتخصیص", "هزینه جاری تحمیل می‌کند ولی اثر آتی ندارد", "ترک"),
]
for i, (a, n, d) in enumerate(rows, start=1):
    cell(i, 0, a, 12, False, INK, CARD)
    cell(i, 1, n, 12, False, INK, CARD)
    cell(i, 2, d, 16, True, TEAL, CARD)
# thick ink rule under the header row (axes are rules, not fills §6)
add_rect(s, tx, ty + Inches(1.0), tw, Pt(3), INK)
notes(s, "جدول محور: سطرها دسته‌ها و ستون‌ها دیدگاه‌ها هستند. ستون راست محور است. "
         "پایان هر سطح یکدست است.")

# ----------------------------------------------------------------------------
# S4/p3 — Premise → conclusion (criterion first: the process needs it)
# ----------------------------------------------------------------------------
s = prs.slides.add_slide(blank)
bg(s)
header(s, "ارزیابی بی‌معیار تشریفاتی می‌شود", "مقدمه و نتیجه", num=3)
ry, rh = Inches(2.0), Inches(3.6)
# RIGHT: premise
add_rect(s, Inches(6.9), ry, Inches(5.83), rh, CARD)
_, rtf = textbox(s, Inches(7.25), ry + Inches(0.3), Inches(5.13), rh - Inches(0.6))
para(rtf, "وضعیت ارزیابی امروز", size=16, bold=True, color=TEAL, first=True,
     space_after=10, bullet=False)
para(rtf, "هر پروژه بدون معیار روشن ادامه می‌یابد", size=14, color=INK,
     space_after=8, bullet=True)
para(rtf, "داوری سلیقه‌ای به حفظ وضع موجود می‌انجامد", size=14, color=INK,
     space_after=0, bullet=True)
# LEFT: consequence (standalone noun-phrase heading)
add_rect(s, MARGIN, ry, Inches(5.7), rh, CARD)
_, ltf = textbox(s, MARGIN + Inches(0.35), ry + Inches(0.3), Inches(5.0), rh - Inches(0.6))
para(ltf, "تصمیم قابل دفاع و تکرارپذیر", size=16, bold=True, color=TEAL, first=True,
     space_after=10, bullet=False)
para(ltf, "معیار توقف پیش از شروع ثبت می‌شود", size=14, color=INK,
     space_after=8, bullet=True)
para(ltf, "همان معیار در داوری اعمال می‌شود", size=14, color=INK,
     space_after=0, bullet=True)
# single center triangle pointing LEFT (right→left flow §4.27)
tri = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(6.4),
                         Inches(3.55), Inches(0.5), Inches(0.5))
tri.fill.solid()
tri.fill.fore_color.rgb = ACCENT
tri.line.fill.background()
tri.rotation = 270.0
notes(s, "جفت مقدمه و نتیجه. مثلث میانی جهت راست‌به‌چپ را نشان می‌دهد.")

# ----------------------------------------------------------------------------
# S5/p4 — Process steps RTL (step 1 rightmost; pivotal stage recolored + legend)
# ----------------------------------------------------------------------------
s = prs.slides.add_slide(blank)
bg(s)
header(s, "ترک موفق نتیجه فرایند است، نه تصمیم لحظه‌ای", "فرایند · راست‌به‌چپ", num=4)
steps = [
    ("۱", "شناسایی", "فهرست فعالیت‌های جاری"),
    ("۲", "ارزیابی", "سنجش با معیار ثبت‌شده"),
    ("۳", "تصمیم", "اعلام رسمی توقف"),
    ("۴", "بازتخصیص", "انتقال منابع"),
]
n = 4
gap = Inches(0.35)
cw = (SLIDE_W - 2 * MARGIN - (n - 1) * gap) / n
cy, ch = Inches(2.2), Inches(3.3)
for i, (num_, name_, desc_) in enumerate(steps):
    x = SLIDE_W - MARGIN - (i + 1) * cw - i * gap
    pivotal = (num_ == "۳")
    add_rect(s, x, cy, cw, ch, ACCENT if pivotal else CARD)
    _, tf = textbox(s, x + Inches(0.25), cy + Inches(0.35), cw - Inches(0.5), ch - Inches(0.7))
    ink = INK  # dark-on-orange for the pivotal stage (white fails contrast)
    sub = SECONDARY
    para(tf, num_, size=30, bold=True, color=ink, align=PP_ALIGN.CENTER,
         first=True, space_after=4, bullet=False)
    para(tf, name_, size=17, bold=True, color=ink, align=PP_ALIGN.CENTER,
         space_after=4, bullet=False)
    para(tf, desc_, size=13, color=sub, align=PP_ALIGN.CENTER, space_after=0,
         bullet=False)
    if i < n - 1:
        # connector pointing LEFT (flow direction right→left) in the gutter
        gx = x - gap / 2 - Inches(0.2)
        tri = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, gx,
                                 cy + ch / 2 - Inches(0.22), Inches(0.4), Inches(0.44))
        tri.fill.solid()
        tri.fill.fore_color.rgb = SECONDARY
        tri.line.fill.background()
        tri.rotation = 270.0
# legend for the single recolored stage (§5.2)
lx = Inches(7.6)
add_rect(s, lx, Inches(5.95), Inches(0.35), Inches(0.35), ACCENT)
_, ltf = textbox(s, lx - Inches(2.6), Inches(5.92), Inches(2.5), Inches(0.42))
para(ltf, "نقطه عطف: تصمیم (گام ۳)", size=12, color=SECONDARY, first=True)
notes(s, "فرایند چهارگامی از راست به چپ خوانده می‌شود. گام سوم نقطه عطف است.")

# ----------------------------------------------------------------------------
# S6/p5 — Chart + insight (So-what in title; educational numbers)
# ----------------------------------------------------------------------------
s = prs.slides.add_slide(blank)
bg(s)
header(s, "در مثال آموزشی، گروه جسورانه بالاتر است", "تک‌نمودار و پیام", num=5)
data = CategoryChartData()
# Deliberate RTL order: baseline (ایستا) rightmost, climax (جسورانه) leftmost —
# a Persian reader moves right→left and the bars rise along the reading path.
data.categories = ["بازتخصیص جسورانه", "بازتخصیص میانه", "بدون بازتخصیص"]
data.add_series("شاخص بازده", (12, 8, 5))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4.83),
                        Inches(1.9), Inches(7.9), Inches(4.6), data)
chart = gf.chart
set_chart_font(chart)
chart.has_legend = False
chart.has_title = True
chart.chart_title.text_frame.text = "شاخص بازده آموزشی (امتیاز) در سه سطح"
plot = chart.plots[0]
plot.gap_width = 90
ser = plot.series[0]
ser.points[0].format.fill.solid()
ser.points[0].format.fill.fore_color.rgb = ACCENT
for k in (1, 2):
    ser.points[k].format.fill.solid()
    ser.points[k].format.fill.fore_color.rgb = SECONDARY
plot.has_data_labels = True
labels = plot.data_labels
labels.show_value = True
labels.position = XL_LABEL_POSITION.OUTSIDE_END
# LEFT: insight
add_rect(s, MARGIN, Inches(1.9), Inches(3.9), Inches(4.6), CARD)
_, itf = textbox(s, MARGIN + Inches(0.35), Inches(2.2), Inches(3.2), Inches(4.0))
para(itf, "پیام نمودار", size=15, bold=True, color=TEAL, first=True,
     space_after=8, bullet=False)
para(itf, "این نمودار واقعی نیست؛ الگوی موردانتظار را نشان می‌دهد.", size=14,
     color=INK, line=1.25, space_after=10, bullet=False)
para(itf, "ستون نارنجی گروه جسورانه است که در این مثال فرضی بالاتر است.", size=13, color=SECONDARY,
     space_after=0, bullet=False)
_, srctf = textbox(s, MARGIN, Inches(6.72), Inches(6.0), Inches(0.28))
para(srctf, "منبع: اعداد فرضی و آموزشی", size=11, color=FAINT, first=True)
notes(s, "نمودار آموزشی با اعداد فرضی. ترتیب دسته‌ها راست‌به‌چپ است: پایه در راست "
         "و گروه جسورانه در چپ. پیام: جسارت در بازتخصیص با بازده همراه است. "
         "قلم نمودار وزیرمتن است.")

# ----------------------------------------------------------------------------
# S7/p6 — Issue / solution (چالش RIGHT, راه‌حل LEFT with row axis)
# ----------------------------------------------------------------------------
s = prs.slides.add_slide(blank)
bg(s)
header(s, "ظرفیت آزاد با توقف به دست می‌آید", "چالش و راه‌حل", num=6)
cy, ch = Inches(1.9), Inches(4.6)
w2 = (SLIDE_W - 2 * MARGIN - Inches(0.33)) / 2
# RIGHT: challenge
add_rect(s, SLIDE_W - MARGIN - w2, cy, w2, ch, CARD)
_, rtf = textbox(s, SLIDE_W - MARGIN - w2 + Inches(0.35), cy + Inches(0.3),
                 w2 - Inches(0.7), ch - Inches(0.6))
para(rtf, "چالش", size=16, bold=True, color=TEAL, first=True, space_after=10,
     bullet=False)
for line_ in ("جلسه هفتگی بدون تصمیم برگزار می‌شود",
              "گزارش تکراری کم‌خواننده تولید می‌شود",
              "پروژه کم‌اثر هر سال تمدید می‌شود",
              "انرژی تیم برای کار کم‌اثر مصرف می‌شود"):
    para(rtf, line_, size=15, color=INK, space_after=8, bullet=True)
# LEFT: solution with row axis (one-off «:» cuts §7.12)
add_rect(s, MARGIN, cy, w2, ch, CARD)
_, ltf = textbox(s, MARGIN + Inches(0.35), cy + Inches(0.3), w2 - Inches(0.7),
                 ch - Inches(0.6))
para(ltf, "راه‌حل", size=16, bold=True, color=TEAL, first=True, space_after=10,
     bullet=False)
for line_ in ("لغو: جلسه هفتگی بدون تصمیم",
              "جایگزینی: یک صفحه به‌جای گزارش تکراری",
              "آزادسازی: بودجه پروژه متوقف‌شده",
              "اثر: تمرکز تیم بر یک کار اثربخش"):
    para(ltf, line_, size=15, color=INK, space_after=8, bullet=True)
notes(s, "نمونه عملی: چهار اقدام توقف و اثر هر یک. ستون چپ محور سطری دارد. "
         "ردیف چهارم زمینه است، نه مورد توقف.")
_, stf = textbox(s, MARGIN, Inches(6.72), Inches(5.0), Inches(0.28))
para(stf, "مثال آموزشی", size=11, color=FAINT, first=True)

# ----------------------------------------------------------------------------
# S8/p7 — Decision page
# ----------------------------------------------------------------------------
s = prs.slides.add_slide(blank)
bg(s)
header(s, "امسال کدام کارها را متوقف می‌کنیم؟", "تصمیم", num=7)
cols = [
    ("تصمیم موردنظر", ["لغو جلسه بدون تصمیم", "توقف گزارش تکراری و جایگزینی یک‌صفحه‌ای",
                       "توقف پروژه کم‌اثر"]),
    ("پیش‌فرض‌ها", ["ملاک لازم‌بودن جلسه روشن است", "معیار سنجش اثر گزارش روشن است",
                    "مقصد بودجه آزادشده مشخص است"]),
    ("درخواست‌ها", ["در جلسه امروز رأی‌گیری شود", "مجری پیگیری توقف معرفی شود"]),
]
cw = (SLIDE_W - 2 * MARGIN - 2 * Inches(0.3)) / 3
cy, ch = Inches(1.9), Inches(4.6)
for i, (head_, lines_) in enumerate(cols):
    x = SLIDE_W - MARGIN - (i + 1) * cw - i * Inches(0.3)
    # card + RIGHT strip first (color behind text), then two stacked boxes
    # (heading / bullets) with no overlap
    add_rect(s, x, cy, cw, ch, CARD)
    sw = Inches(0.16)
    add_rect(s, x + cw - sw, cy + Inches(0.12), sw, ch - Inches(0.24), TEAL)
    _, htf = textbox(s, x + Inches(0.35), cy + Inches(0.22), cw - Inches(0.7),
                     Inches(0.9))
    para(htf, head_, size=16, bold=True, color=INK, first=True, space_after=0,
         bullet=False)
    _, btf = textbox(s, x + Inches(0.35), cy + Inches(1.2), cw - Inches(0.7),
                     ch - Inches(1.55))
    for j, line_ in enumerate(lines_):
        para(btf, line_, size=13.5, color=INK, first=(j == 0), space_after=8,
             bullet=True)
notes(s, "صفحه تصمیم: چه چیزی، بر چه مبنایی، چه کمکی لازم است. "
         "رأی فقط برای سه توقف است.")

# ----------------------------------------------------------------------------
# S9 — Back cover (no message, unnumbered §4.24)
# ----------------------------------------------------------------------------
s = prs.slides.add_slide(blank)
bg(s)
_, tf = textbox(s, MARGIN, Inches(2.1), SLIDE_W - 2 * MARGIN, Inches(3.4),
                anchor=MSO_ANCHOR.MIDDLE)
para(tf, "گروه برنامه‌ریزی استراتژیک", size=30, bold=True, color=INK, first=True,
     align=PP_ALIGN.CENTER, space_after=10, bullet=False)
para(tf, "ارائه آموزشی ترک استراتژیک", size=16, color=SECONDARY,
     align=PP_ALIGN.CENTER, space_after=10, bullet=False)
para(tf, "راه تماس: مجری پیگیری توقف", size=14, color=SECONDARY,
     align=PP_ALIGN.CENTER, bullet=False)
_, ctf = textbox(s, MARGIN, Inches(7.12), Inches(6.0), Inches(0.3))
para(ctf, "روش برگرفته از consulting-pptx-skill", size=10, color=FAINT,
     align=PP_ALIGN.LEFT, rtl=True, first=True)
notes(s, "پشت‌جلد. بدون پیام. فقط نام و تماس.")

prs.save(OUT)
print(f"saved {OUT} with {len(prs.slides._sldIdLst)} slides")
