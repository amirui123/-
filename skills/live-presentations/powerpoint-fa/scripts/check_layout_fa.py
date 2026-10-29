# -*- coding: utf-8 -*-
"""Static geometry audit for Persian .pptx decks (no renderer needed).

Port of consulting-pptx-skill `scripts/check_layout.mjs` issue taxonomy
(footer overlap, right/bottom overflow) to native python-pptx geometry —
the Playwright/Node dependency is intentionally not ported.

Usage:
    python3 scripts/check_layout_fa.py deck.pptx

Render audit (silent-clip catcher) stays manual per SKILL.md procedure:
    soffice --headless --convert-to pdf --outdir render deck.pptx
    pdftoppm -png render/deck.pdf render/slide
"""
import argparse
import sys
from pathlib import Path

FOOTER_TOP = 7.06 * 914400
SLIDE_W = 13.333 * 914400
SLIDE_H = 7.5 * 914400
TOL = 1500  # ~0.0016in tolerance


def rect(sh):
    try:
        if sh.left is None or sh.top is None or sh.width is None or sh.height is None:
            return None
        return (sh.left, sh.top, sh.left + sh.width, sh.top + sh.height)
    except Exception:
        return None


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("path", help="deck.pptx")
    a = ap.parse_args()
    p = a.path
    if not Path(p).exists():
        sys.exit(f"not found: {p}")
    try:
        from pptx import Presentation
    except ImportError:
        sys.exit("python-pptx is required: pip install python-pptx==1.0.2")
    prs = Presentation(p)
    issues = []
    for i, slide in enumerate(prs.slides, 1):
        boxes = []
        for sh in slide.shapes:
            r = rect(sh)
            if r is None:
                continue
            x0, y0, x1, y1 = r
            name = ""
            try:
                name = sh.text_frame.text.strip().replace("\n", " ")[:30] if sh.has_text_frame else sh.name[:30]
            except Exception:
                pass
            # background full-slide rects are not content
            try:
                if x0 <= TOL and y0 <= TOL and x1 >= SLIDE_W - TOL and y1 >= SLIDE_H - TOL:
                    continue
            except Exception:
                pass
            if x1 > SLIDE_W + TOL:
                issues.append(f"P{i}: right overflow +{(x1 - SLIDE_W) / 914400:.2f}in: «{name}»")
            if y1 > SLIDE_H + TOL:
                issues.append(f"P{i}: bottom overflow +{(y1 - SLIDE_H) / 914400:.2f}in: «{name}»")
            # footer-band overlap (footer rule + footer text themselves excluded)
            try:
                h = sh.height or 0
                w = sh.width or 0
                hairline = h < 914400 * 0.02 and w > 914400 * 11
                in_footer = y0 >= FOOTER_TOP - TOL
            except Exception:
                hairline, in_footer = False, False
            if not hairline and not in_footer and y1 > FOOTER_TOP + TOL and y0 < FOOTER_TOP:
                issues.append(f"P{i}: footer overlap: «{name}» crosses y=7.06in")
            # Only non-empty text boxes participate in pairwise overlap:
            # filled rects/panels behind text (skill pattern: color sits behind
            # the text) carry empty text frames and must not flag.
            try:
                txt = sh.text_frame.text.strip() if sh.has_text_frame else ""
            except Exception:
                txt = ""
            if sh.has_text_frame and txt and (sh.width or 0) > 0 and (sh.height or 0) > 0:
                try:
                    top_in = sh.top if sh.top is not None else 0
                except Exception:
                    top_in = 0
                boxes.append((txt[:30], r, top_in >= FOOTER_TOP - TOL))
        # pairwise content overlap (text boxes only, small intersections ignored;
        # footer-band co-location — brand line + page number sharing the footer
        # row by design — is excluded)
        for a_i in range(len(boxes)):
            for b_i in range(a_i + 1, len(boxes)):
                na, (ax0, ay0, ax1, ay1), a_f = boxes[a_i]
                nb, (bx0, by0, bx1, by1), b_f = boxes[b_i]
                if a_f and b_f:
                    continue
                ix = min(ax1, bx1) - max(ax0, bx0)
                iy = min(ay1, by1) - max(ay0, by0)
                if ix > 5000 and iy > 5000:
                    area = ix * iy / (914400.0 ** 2)
                    if area > 0.02:
                        issues.append(f"P{i}: overlap ~{area:.2f}in²: «{na}» × «{nb}»")
                        break
            else:
                continue
            break
    seen, uniq = set(), []
    for m in issues:
        if m not in seen:
            seen.add(m)
            uniq.append(m)
    if uniq:
        for m in uniq[:20]:
            print("FAIL  " + m)
        print(f"\n{len(uniq)} layout FAIL")
        sys.exit(1)
    print("layout OK (footer overlap / overflow / pairwise text overlap)")


if __name__ == "__main__":
    main()
