# -*- coding: utf-8 -*-
"""Mechanical rule check for Persian python-pptx decks (FAIL 0 to ship).

Port of consulting-pptx-skill `scripts/check_deck.py` (MIT) to the
`powerpoint-fa` workflow, merged with the RTL checks from `validate.py`.

Usage:
    python3 scripts/check_deck_fa.py deck.pptx
    python3 scripts/check_deck_fa.py assets/SuperTemplate-fa.pptx --template

Canon: references/slide-rules-fa.md
Exit code: 1 when any FAIL remains.
Requires: python-pptx (stdlib + python-pptx only).
"""
import argparse
import re
import sys
import zipfile
from pathlib import Path

# Brand-banned legacy colors go here (HEX 6 digits, no '#'). Empty = no check.
OLD_COLORS: list = []
TITLE_MAX = 40
EMU_W, EMU_H = 12192000, 6858000

fails: list = []
warns: list = []


def fail(msg):
    fails.append(msg)


def warn(msg):
    warns.append(msg)


TEMPLATE_MODE = False

# --- Persian term-drift pairs (slide-rules-fa §7.6: 1 doc 1 term) -----------
# Both spellings in one deck -> WARN (distinct-meaning uses are eye-waived).
TERM_VARIANTS = [
    ("ی فارسی", r"ی", "ي عربی", r"ي"),
    ("ک فارسی", r"ک", "ك عربی", r"ك"),
    ("ارقام فارسی", r"[۰-۹]", "ارقام لاتین", r"(?<![A-Za-z0-9:/-])[0-9]"),
]

# --- Persian AI-smell words (slide-rules-fa §7.9 / ai-smell-fa.md) -----------
# High-precision subset only -> WARN (context-justified hits eye-waived).
AI_SMELL_WORDS = [
    "بسیار", "فوق‌العاده", "بی‌نهایت", "انقلابی", "نسل بعدی",
    "اغراق نیست", "چیزی نیست جز",
    "در این راستا", "با توجه به", "در عصر حاضر", "در دنیای امروز",
    "همان‌طور که مستحضرید", "به شرح ذیل", "خواهشمند است",
    "از منظر", "نکته کلیدی",
    "سینرژی", "سرتاسری", "ارزش افزوده",
]

PLACEHOLDER_RE = re.compile(
    r"متن نمونه|عنوان نمونه|نام شرکت|نام سند|شرکت نمونه"
    r"|Text\s*\d|ラベル\s*\d|タイトル\s*\d|Source\s*\d|YYYY|XXX+|TODO|TBD|Lorem"
    r"|[◯○]{2,}|^۱۲۳$|^شرکت$|^عنوان$"
)


def check_ai_smell(pages):
    hits = {}
    for i, txt in pages:
        found = [w for w in AI_SMELL_WORDS if w in txt]
        if found:
            hits[i] = found
    if hits:
        detail = "،".join(f"P{i}:«{'/'.join(ws[:3])}»" for i, ws in sorted(hits.items())[:6])
        warn(f"AI-smell words: {detail} (§7.9 / ai-smell-fa.md — plain verbs first)")
    dash_pages = sorted({i for i, txt in pages if " — " in txt or " —" in txt})
    if dash_pages:
        warn(f"Dash « — » joins on P{dash_pages} (typical AI tic — use period/`:`/parens. §7.9)")


def check_terms(pages):
    for la, pa, lb, pb in TERM_VARIANTS:
        try:
            ha = sorted({i for i, t in pages if re.search(pa, t)})
            hb = sorted({i for i, t in pages if re.search(pb, t)})
        except re.error:
            continue
        if ha and hb:
            warn(f"Term drift: «{la}» P{ha} vs «{lb}» P{hb} (§7.6 one-doc-one-term; distinct meanings eye-waived)")


def fw_len(t):
    """Full-width-equivalent length for Persian decks (see fa_helpers)."""
    total = 0.0
    for ch in t:
        o = ord(ch)
        if o in (0x200C, 0x200D):
            continue
        if o >= 0x3000 or 0x0600 <= o <= 0x06FF or 0xFB50 <= o <= 0xFEFF:
            total += 1.0
        else:
            total += 0.5
    return total


def check_body_placeholders(idx, text):
    """Draft body placeholders must not survive into delivery decks."""
    if not TEMPLATE_MODE and text and PLACEHOLDER_RE.search(text):
        m = PLACEHOLDER_RE.search(text)
        fail(f"P{idx}: body placeholder «{m.group(0)}» left (replace with storyline — §2.8)")


def check_title(idx, title, explicit_break=False):
    t = title.strip()
    if not t:
        warn(f"P{idx}: empty title (ok for cover/door only)")
        return
    if re.search(r"(لطفاً.*بفرمایید|خواهشمند است.*)$", t):
        fail(f"P{idx}: title uses polite padding: «{t}»")
    tlen = fw_len(t)
    if tlen > TITLE_MAX * 2:
        fail(f"P{idx}: title {tlen:.0f} full-equiv chars (>{TITLE_MAX*2}, won't fit 2 lines — narrow the claim): «{t}»")
    elif tlen > TITLE_MAX and not explicit_break:
        warn(f"P{idx}: title {tlen:.0f} full-equiv chars (expect 2 lines — break at meaning, no orphan, never shrink-to-fit): «{t}»")
    if re.match(r"^(Step|STEP|مرحله|فاز|گام)\s*\d", t):
        fail(f"P{idx}: title concatenates Step tag (use corner chip): «{t}»")
    if re.match(r"^(این|آن|همان|تا اینجا)", t):
        fail(f"P{idx}: title starts with cross-slide reference: «{t}»")
    if t.count("（") + t.count("(") >= 2:
        warn(f"P{idx}: many parens in title: «{t}»")
    if not TEMPLATE_MODE and PLACEHOLDER_RE.search(t):
        fail(f"P{idx}: template placeholder left (write titles from storyline — §2.8): «{t}»")


def check_title_variety(titles):
    real = [t.strip() for t in titles if t and t.strip()]
    if len(real) < 5:
        return
    molds = {
        "«… است» endings": lambda t: t.endswith("است"),
        "count-led («۳ …»)": lambda t: re.match(r"^[0-9۰-۹]+", t),
        "«label:» prefixes": lambda t: re.match(r"^[^:：]{1,12}[:：]", t),
    }
    for name, f in molds.items():
        try:
            hits = [t for t in real if f(t)]
        except Exception:
            continue
        if len(hits) >= max(4, int(len(real) * 0.6)):
            warn(f"Title molds cluster on {name} ({len(hits)}/{len(real)}). Vary per claim — §2.8")


# --- number match (slide-rules-fa §2.9) --------------------------------------
FA_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
COUNT_WORDS = r"(?:مرحله|مراحل|گام|فاز|ستون|محور|بخش|دلیل|راهکار|قدم|Phase|Step|STEP)"
ENUM_PAT = r"(?:مرحله|گام|فاز|قدم|STEP|Step|PHASE|Phase)"


def _to_int(s):
    s = s.translate(FA_DIGITS)
    return int(s) if s.isdigit() else None


def check_count_match(idx, title, text):
    if not title or not text:
        return
    for m in re.finditer(r"([0-9۰-۹٠-٩]+)\s*(" + COUNT_WORDS + r")", title):
        n = _to_int(m.group(1))
        if not n or not (2 <= n <= 12):
            continue
        found = [v for v in (_to_int(x.group(1)) for x in re.finditer(ENUM_PAT + r"\s*([0-9۰-۹٠-٩]+)", text)) if v]
        if found and max(found) != n:
            fail(f"P{idx}: title «{m.group(0)}» vs body ordinals (max {max(found)}) mismatch (§2.9): «{title}»")


# --- geometry estimators (§2.13 orphan / §5.13 fill-rate) ---------------------
def check_orphan(idx, shape, label="title"):
    try:
        if not (shape.has_text_frame and shape.width):
            return
        tf = shape.text_frame
        raw = tf.text
        if not raw.strip():
            return
        szs = [r.font.size.pt for p in tf.paragraphs for r in p.runs if r.font.size]
        if not szs:
            return
        pt = max(szs)
        w_in = shape.width / 914400.0
        cap = max(4, int(w_in / (pt / 72.0)))
        for line in raw.split("\n"):
            L = fw_len(line.strip())
            if L <= cap:
                continue
            rows = int(-(-L // cap))
            last = L - cap * (rows - 1)
            if 0 < last <= 3:
                warn(f"P{idx}: {label} orphan-risk (line ~{last:.0f} chars stranded — break at meaning. §2.13): «{line.strip()[:40]}»")
    except Exception:
        return


def check_fill_ratio(idx, slide, title_shape):
    try:
        top_lim = 1.6 * 914400
        bot_lim = 6.8 * 914400
        tops, bots = [], []
        for sh in slide.shapes:
            if sh is title_shape or sh.top is None or sh.height is None:
                continue
            t, b = sh.top, sh.top + sh.height
            if b <= top_lim or t >= bot_lim:
                continue
            if sh.height < 914400 * 0.02 and sh.width and sh.width > 914400 * 11:
                continue  # full-width hairlines (title/footer rules)
            tops.append(max(t, top_lim))
            bots.append(min(b, bot_lim))
        if not tops:
            return
        cov = (max(bots) - min(tops)) / (bot_lim - top_lim)
        if cov < 0.55:
            warn(f"P{idx}: fill-rate {cov*100:.0f}% (body < half the page — add info or change type; never stretch/decor-fill. §5.13)")
    except Exception:
        return


PERSIAN_RE = re.compile(r"[\u0600-\u06FF]")


def check_pptx(path, strict=False):
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        sys.exit("python-pptx is required: pip install python-pptx==1.0.2")
    prs = Presentation(path)
    if abs(prs.slide_width - EMU_W) > 2000 or abs(prs.slide_height - EMU_H) > 2000:
        fail(f"Slide size {prs.slide_width}x{prs.slide_height} != 16:9 {EMU_W}x{EMU_H}")

    titles = []
    term_pages = []
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        for n in sorted(names, key=lambda s: int(re.search(r"slide(\d+)", s).group(1))):
            xml = z.read(n).decode("utf8", "ignore")
            idx = int(re.search(r"slide(\d+)", n).group(1))
            term_pages.append((idx, " ".join(re.findall(r"<a:t>(.*?)</a:t>", xml, re.S))))
            rr = 0
            for m in re.finditer(r"<p:sp>.*?</p:sp>", xml, re.S):
                sp = m.group(0)
                if 'prst="roundRect"' in sp:
                    h = re.search(r'<a:ext cx="\d+" cy="(\d+)"', sp)
                    if h and int(h.group(1)) >= 365760:
                        rr += 1
            if rr:
                msg = f"P{idx}: large roundRect boxes x{rr} (rect default; pills <0.4in excepted — §5.1)"
                (fail if strict else warn)(msg)
            fb = 0
            for m in re.finditer(r"<p:sp>.*?</p:sp>", xml, re.S):
                sp = m.group(0)
                spPr = re.search(r"<p:spPr>.*?</p:spPr>", sp, re.S)
                if not spPr:
                    continue
                pr = spPr.group(0)
                h = re.search(r'<a:ext cx="\d+" cy="(\d+)"', pr)
                if not (h and int(h.group(1)) >= 365760):
                    continue
                ln = re.search(r"<a:ln[ >].*?</a:ln>", pr, re.S)
                filled = "<a:solidFill>" in re.sub(r"<a:ln[ >].*?</a:ln>", "", pr, flags=re.S)
                if filled and ln and "<a:solidFill>" in ln.group(0):
                    fb += 1
            if fb:
                warn(f"P{idx}: filled shapes with border x{fb} (filled cards carry no line — §5.3)")
            for c in OLD_COLORS:
                if f'val="{c}"' in xml or f'val="{c.lower()}"' in xml:
                    fail(f"P{idx}: banned color {c}")
            lb = len(re.findall(r"<a:t>\s*•", xml))
            if lb:
                fail(f"P{idx}: typed «•» bullets in run text x{lb} (use buChar/buNone — §7.3)")
            ld = len(re.findall(r"<a:t>\s*[–‐-]\s\s", xml))
            if ld:
                warn(f"P{idx}: typed «– » L2-marker-like text x{ld} (use buChar — §7.3)")

    # per-slide semantic + RTL checks (merged validate.py coverage)
    rtl_paras = 0
    fa_runs = 0
    persian_paras = 0
    persian_no_rtl = 0
    cs_missing = 0
    pics_missing_alt = []
    notes_pages = 0
    has_chart = has_table = False

    for i, s in enumerate(prs.slides, 1):
        title = ""
        if s.shapes.title is not None and s.shapes.title.has_text_frame:
            title = s.shapes.title.text_frame.text
        else:
            cands = []
            for sh in s.shapes:
                if sh.has_text_frame and sh.top is not None and sh.top < Emu(1097280):
                    try:
                        sz = max((r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size), default=0)
                    except Exception:
                        sz = 0
                    cands.append((sz, sh.text_frame.text))
            if cands:
                title = max(cands)[1]
        title_shape = None
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text == title:
                title_shape = sh
                break
        if title_shape is not None:
            check_orphan(i, title_shape, "title")
        if i == 1:
            for sh in s.shapes:
                if sh is not title_shape and sh.has_text_frame:
                    try:
                        szs = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
                    except Exception:
                        szs = []
                    if szs and max(szs) >= 13 and len(sh.text_frame.text) < 80:
                        check_orphan(i, sh, "cover text")
        else:
            check_fill_ratio(i, s, title_shape)
        explicit_break = "\n" in title.strip() and all(fw_len(x.strip()) <= TITLE_MAX for x in title.split("\n"))
        body_text = dict(term_pages).get(i, "")
        check_title(i, title.replace("\n", " "), explicit_break)
        check_count_match(i, title, body_text)
        # body placeholders (draft leftovers) block delivery like title ones
        body_wo_title = body_text.replace(title.strip(), "", 1) if title.strip() else body_text
        if body_wo_title.strip():
            check_body_placeholders(i, body_wo_title)
        titles.append(title.replace("\n", " "))
        # subtitle suspicion: thin single line right under title band
        for sh in s.shapes:
            if sh.has_text_frame and sh.top is not None and Emu(1097280) <= sh.top < Emu(1600200):
                txt = sh.text_frame.text.strip()
                if txt and "\n" not in txt and len(txt) < 60 and txt != title:
                    try:
                        szs = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
                        bolds = [bool(r.font.bold) for p in sh.text_frame.paragraphs for r in p.runs]
                    except Exception:
                        szs, bolds = [], []
                    if szs and max(szs) <= 12 and not any(bolds):
                        warn(f"P{i}: subtitle-like line under title (banned — §3): «{txt}»")
        # shapes: RTL + fonts + media + tables
        for shape in s.shapes:
            if getattr(shape, "shape_type", None) == 13:  # PICTURE
                try:
                    nv = shape._element.nvPicPr
                    nm = nv.cNvPr.get("name")
                    ds = nv.cNvPr.get("descr")
                except Exception:
                    nm, ds = "(pic)", ""
                if not ds:
                    pics_missing_alt.append((i, nm))
            if getattr(shape, "has_chart", False):
                has_chart = True
            if getattr(shape, "has_table", False):
                has_table = True
                try:
                    t = shape.table
                    hdr_sizes, body_sizes = [], []
                    for r in range(len(t.rows)):
                        for c in range(len(t.columns)):
                            cell = t.cell(r, c)
                            try:
                                sz = max((run.font.size.pt for p in cell.text_frame.paragraphs for run in p.runs if run.font.size), default=0)
                            except Exception:
                                sz = 0
                            (hdr_sizes if r == 0 else body_sizes).append(sz)
                    if hdr_sizes and body_sizes and max(hdr_sizes) and max(body_sizes):
                        if max(hdr_sizes) < max(body_sizes) + 2 - 1e-6:
                            fail(f"P{i}: table header {max(hdr_sizes):.1f}pt < body {max(body_sizes):.1f}pt +2 (§6)")
                        elif max(hdr_sizes) < max(body_sizes) + 3 - 1e-6:
                            warn(f"P{i}: table header {max(hdr_sizes):.1f}pt < body +3pt guide (§6)")
                except Exception:
                    pass
            if shape.has_text_frame:
                from pptx.oxml.ns import qn
                for p in shape.text_frame.paragraphs:
                    txt = "".join(r.text for r in p.runs)
                    has_fa = bool(PERSIAN_RE.search(txt))
                    try:
                        pPr = p._p.find(qn("a:pPr"))
                        rtl = pPr is not None and pPr.get("rtl") == "1"
                    except Exception:
                        rtl = False
                    if has_fa:
                        persian_paras += 1
                        if rtl:
                            rtl_paras += 1
                        else:
                            persian_no_rtl += 1
                    elif rtl:
                        rtl_paras += 1
                    for r in p.runs:
                        try:
                            rPr = r._r.find(qn("a:rPr"))
                        except Exception:
                            rPr = None
                        if rPr is None:
                            continue
                        try:
                            from pptx.oxml.ns import qn as _qn
                            lang = rPr.get("lang")
                            if lang and lang.upper().startswith("FA"):
                                fa_runs += 1
                            cs = rPr.find(_qn("a:cs"))
                            lat = rPr.find(_qn("a:latin"))
                            if has_fa or (lang and lang.upper().startswith("FA")):
                                if cs is None:
                                    cs_missing += 1
                                elif lat is not None and cs.get("typeface") != lat.get("typeface"):
                                    cs_missing += 1
                        except Exception:
                            pass
        try:
            if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
                notes_pages += 1
        except Exception:
            pass

    # Latin-caps decor kicker (Persian decks use Persian tag chips; §7.22)
    try:
        body_texts = [t for _, t in term_pages]
        caps = []
        for t in body_texts:
            caps += re.findall(r"(?:\d{2}\s*[・·/|]\s*)?[A-Z][A-Z0-9 &/·・\-]{5,}", t)
        caps = [c for c in caps if re.search(r"[A-Z]{3,}\s+[A-Z]{2,}", c)]
        if caps:
            warn(f"Latin-caps decor kicker-like strings x{len(caps)} (§7.22 — Persian tag chips instead)")
    except Exception:
        pass

    check_terms(term_pages)
    check_ai_smell(term_pages)
    if not TEMPLATE_MODE:
        check_title_variety(titles)

    # RTL summary gates
    print("=== RTL summary ===")
    print(f"Persian paragraphs: {persian_paras} (rtl=1: {rtl_paras}, missing rtl: {persian_no_rtl})")
    print(f"FARSI-tagged runs: {fa_runs}")
    print(f"Slides with notes: {notes_pages}/{len(list(prs.slides))}  chart: {has_chart}  table: {has_table}")
    if persian_paras and rtl_paras == 0:
        fail("No rtl=\"1\" Persian paragraphs (alignment is not direction — §3 procedure)")
    elif persian_no_rtl:
        warn(f"{persian_no_rtl} Persian paragraphs lack rtl=\"1\" (§3 procedure)")
    if cs_missing:
        warn(f"{cs_missing} Persian runs lack matching a:cs typeface (set latin+ea+cs — Pitfalls)")
    for i, nm in pics_missing_alt:
        fail(f"P{i}: picture «{nm}» lacks alt text")
    if notes_pages == 0:
        warn("No speaker notes found (notes carry cues/sources/alt drafts — §5 procedure)")
    return titles


def main():
    global TEMPLATE_MODE
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("path", help="deck.pptx to check")
    ap.add_argument("--template", action="store_true", help="checking a pattern book itself (placeholders allowed)")
    ap.add_argument("--strict", action="store_true", help="large roundRect becomes FAIL instead of WARN")
    a = ap.parse_args()
    TEMPLATE_MODE = a.template
    p = a.path
    if not p.lower().endswith(".pptx"):
        sys.exit("powerpoint-fa checks .pptx decks (HTML pipeline is out of scope).")
    if not Path(p).exists():
        sys.exit(f"not found: {p}")
    titles = check_pptx(p, strict=a.strict)
    print("\n=== Title list (read top-to-bottom: one story?) ===")
    for i, t in enumerate(titles, 1):
        print(f"{i:>3}  {t or '(empty)'}")
    print()
    for w in warns:
        print("WARN ", w)
    for f in fails:
        print("FAIL ", f)
    print(f"\n{len(fails)} FAIL / {len(warns)} WARN")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
