# -*- coding: utf-8 -*-
"""Assemble a Persian draft deck from pattern IDs (draft first, adjust after).

Port of consulting-pptx-skill `scripts/new_deck.py` concept (MIT) to
native python-pptx: pick patterns, get one themed RTL .pptx draft with
renumbered footers, then replace placeholders with the real storyline.

Usage:
    python3 scripts/pptx_new_fa.py --list
    python3 scripts/pptx_new_fa.py --parts cover,map,axis,chart,decision,back --title "نام سند" -o mydeck.pptx

Catalog: references/archetype-catalog-fa.md
Rules: references/slide-rules-fa.md (never force a slide into a type).
"""
import argparse
import sys

sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parent))

PATTERNS = [
    ("cover", "Cover — title + context + hero visual"),
    ("map", "Whole-map overview (each row with → P.n refs)"),
    ("toc", "Table of contents (one column)"),
    ("door", "Section door (SECTION n / N, unnumbered)"),
    ("chevron", "Process chevron, RTL (step 1 rightmost)"),
    ("vsteps", "Vertical steps + side comment"),
    ("premise", "Premise → conclusion (single center triangle)"),
    ("axis", "Axis table (rows = items, cols = viewpoints)"),
    ("compare", "Choice comparison table"),
    ("risk", "Risk list (risk + countermeasure pairs)"),
    ("stat", "Large-number readout + reading"),
    ("chart", "Single chart + insight (So-what mandatory)"),
    ("bars", "Ranked horizontal bars + note"),
    ("matrix", "2x2 matrix"),
    ("harvey", "Harvey-ball / heatmap status table"),
    ("roadmap", "Roadmap (RTL time right→left)"),
    ("gantt", "Gantt schedule"),
    ("decision", "Decision page (decide / premises / asks)"),
    ("case", "Issue / solution (چالش RIGHT, راه‌حل LEFT)"),
    ("back", "Back cover (logo + contact only)"),
]

NO_NUMBER = {"cover", "door", "back"}


def build(parts, title, out):
    from pptx import Presentation
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.chart.data import CategoryChartData
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
    from fa_helpers import (
        SLIDE_W, SLIDE_H, MARGIN, SAND, INK, SECONDARY, ACCENT, TEAL,
        CARD, CARDBORDER, FAINT, DARKPANEL, FONT, set_run_font, set_para_rtl,
        set_bullet,
    )

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]
    cp = prs.core_properties
    cp.title = title
    cp.author = "powerpoint-fa (scripts/pptx_new_fa.py draft)"
    cp.subject = "Persian RTL draft — replace placeholders with the storyline"

    def bg(slide, color=SAND):
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
        shp.fill.solid()
        shp.fill.fore_color.rgb = color
        shp.line.fill.background()
        shp.shadow.inherit = False

    def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
        return tb, tf

    def para(tf, text, size=18, bold=False, color=INK, first=False,
             rtl=True, align=None, bullet=None, space_after=6):
        from pptx.enum.text import PP_ALIGN as _A
        p = tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()
        p.alignment = align if align is not None else (_A.RIGHT if rtl else _A.LEFT)
        p.space_after = Pt(space_after)
        if rtl:
            set_para_rtl(p, True)
        if bullet == "off":
            set_bullet(p, char=None)
        elif bullet:
            set_bullet(p, level=0, char="•")
        r = p.add_run()
        r.text = text
        set_run_font(r, size=size, bold=bold, color=color)
        return p

    def header(slide, heading, kicker=None, num=None, total=None):
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.14))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT
        bar.line.fill.background()
        _, tf = textbox(slide, MARGIN, Inches(0.42), SLIDE_W - 2 * MARGIN, Inches(1.1))
        if kicker:
            para(tf, kicker, size=13, bold=True, color=TEAL, first=True, space_after=2)
        para(tf, heading, size=30, bold=True, color=INK, space_after=0)
        rule = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, MARGIN, Inches(7.06), SLIDE_W - 2 * MARGIN, Pt(1))
        rule.fill.solid()
        rule.fill.fore_color.rgb = CARDBORDER
        rule.line.fill.background()
        _, ftf = textbox(slide, MARGIN, Inches(7.12), Inches(4.0), Inches(0.3))
        para(ftf, title, size=10, color=FAINT, rtl=True, first=True)
        if num is not None:
            _, ntf = textbox(slide, SLIDE_W - MARGIN - Inches(1.0), Inches(7.12), Inches(1.0), Inches(0.3))
            para(ntf, f"{num} / {total}", size=10, color=FAINT, rtl=False,
                 align=PP_ALIGN.RIGHT, first=True)

    def notes(slide, text):
        slide.notes_slide.notes_text_frame.text = text

    page = 0
    total = sum(1 for p in parts if p not in NO_NUMBER)

    for pat in parts:
        s = prs.slides.add_slide(blank)
        numbered = pat not in NO_NUMBER
        if numbered:
            page += 1
        num = page if numbered else None
        if pat == "cover":
            bg(s)
            panel = s.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, MARGIN, Inches(1.35), Inches(7.5), Inches(4.2))
            panel.fill.solid()
            panel.fill.fore_color.rgb = DARKPANEL
            panel.line.fill.background()
            _, tf = textbox(s, MARGIN + Inches(0.6), Inches(1.75), Inches(6.3), Inches(3.6))
            para(tf, "پیش‌نویس powerpoint-fa", size=15, bold=True, color=ACCENT, first=True, space_after=12)
            para(tf, title, size=40, bold=True, color=SAND, space_after=6)
            para(tf, "ادعا را از خط داستانی بنویسید — عنوان نمونه را پر کنید",
                 size=16, color=SAND, space_after=18)
            para(tf, "جایگزین کنید: زمینه، مخاطب، تصمیم موردنظر", size=14, color=SECONDARY)
            notes(s, "جلد. عنوان = ادعا از خط داستانی؛ بدون پیام پایانی جدا.")
        elif pat == "back":
            bg(s)
            _, tf = textbox(s, MARGIN, Inches(2.5), SLIDE_W - 2 * MARGIN, Inches(2.0),
                            anchor=MSO_ANCHOR.MIDDLE)
            para(tf, "نام شرکت · تماس", size=28, bold=True, color=INK, first=True,
                 align=PP_ALIGN.CENTER)
            para(tf, "پشت‌جلد فقط لوگو + نام + تماس است — بدون پیام (§4.24)",
                 size=14, color=FAINT, align=PP_ALIGN.CENTER)
            notes(s, "پشت‌جلد. بدون پیام.")
        elif pat == "door":
            bg(s)
            _, tf = textbox(s, MARGIN, Inches(2.0), SLIDE_W - 2 * MARGIN, Inches(3.0),
                            anchor=MSO_ANCHOR.MIDDLE)
            para(tf, "SECTION ۱ / ۳", size=16, bold=True, color=TEAL, first=True,
                 align=PP_ALIGN.CENTER, rtl=False)
            para(tf, "عنوان بخش را بنویسید", size=34, bold=True, color=INK,
                 align=PP_ALIGN.CENTER)
            notes(s, "درِ بخش. فقط برای دک‌های بالای ~۱۰ اسلاید؛ شماره صفحه ندارد.")
        elif pat in ("axis", "compare", "risk", "harvey"):
            bg(s)
            heads = {"axis": ("جدول محور", "سطر = مورد، ستون = دیدگاه"),
                     "compare": ("مقایسه گزینه‌ها", "ستون برنده هایلایت + محور برد رقیب"),
                     "risk": ("ریسک‌ها", "هر ریسک با اقدام متقابل"),
                     "harvey": ("وضعیت", "○◐● + — خاکستری برای نامربوط")}[pat]
            header(s, heads[0], kicker=heads[1], num=num, total=total)
            shape = s.shapes.add_table(4, 3, MARGIN, Inches(1.8),
                                       SLIDE_W - 2 * MARGIN, Inches(4.6))
            tbl = shape.table
            hdrs = ["دیدگاه ب", "دیدگاه الف", "مورد / محور"]
            for j, h in enumerate(hdrs):
                cell = tbl.cell(0, j)
                cell.fill.solid()
                cell.fill.fore_color.rgb = INK
                p = cell.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.RIGHT
                set_para_rtl(p, True)
                r = p.add_run()
                r.text = h
                set_run_font(r, size=15, bold=True, color=CARD)
            for r in range(1, 4):
                for c in range(3):
                    cell = tbl.cell(r, c)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = CARD
                    p = cell.text_frame.paragraphs[0]
                    p.alignment = PP_ALIGN.RIGHT
                    set_para_rtl(p, True)
                    run = p.add_run()
                    run.text = "متن نمونه" if (r, c) != (1, 2) else "مورد نمونه"
                    set_run_font(run, size=13)
            notes(s, f"الگوی {pat}. سلول‌ها را با خط داستانی جایگزین کنید؛ قانون جدول §6.")
        elif pat in ("chart", "bars"):
            bg(s)
            header(s, "ادعای نمودار را اینجا بنویسید (So-what اجباری)",
                   kicker="تک‌نمودار + پیام", num=num, total=total)
            data = CategoryChartData()
            data.categories = ["الف", "ب", "ج"]
            data.add_series("مقدار", (8, 5, 3))
            gf = s.shapes.add_chart(
                XL_CHART_TYPE.COLUMN_CLUSTERED, MARGIN, Inches(1.9),
                Inches(7.6), Inches(4.4), data)
            chart = gf.chart
            chart.has_legend = False
            chart.has_title = True
            chart.chart_title.text_frame.text = "عنوان نمودار = نتیجه، نه نام نمودار"
            try:
                from pptx.oxml.ns import qn
                for tag in ("a:latin", "a:ea", "a:cs"):
                    for el in chart._chartSpace.iter(qn(tag)):
                        el.set("typeface", FONT)
            except Exception:
                pass
            _, tf = textbox(s, Inches(8.6), Inches(2.2), Inches(4.1), Inches(3.4))
            para(tf, "پیام نمودار", size=15, bold=True, color=TEAL, first=True, space_after=8)
            para(tf, "متن نمونه: یافته + معنای آن برای تصمیم.", size=14, space_after=12)
            para(tf, "منبع: پایین‌چپ، یک خط", size=11, color=FAINT)
            notes(s, "الگوی نمودار. روند/ترکیب/توزیع فقط نمودار (§5.11)؛ پیام در عنوان.")
        elif pat in ("chevron", "roadmap"):
            bg(s)
            header(s, "فرایند را اینجا بنویسید (گام ۱ در راست‌ترین)",
                   kicker="شه‌پرون راست‌به‌چپ", num=num, total=total)
            n, cw = 4, (SLIDE_W - 2 * MARGIN - 3 * Inches(0.25)) / 4
            for k in range(n):
                k_rev = n - 1 - k  # RTL: step 1 rightmost
                x = MARGIN + k * (cw + Inches(0.25))
                shp = s.shapes.add_shape(MSO_SHAPE.CHEVRON, x, Inches(3.0), cw, Inches(1.6))
                shp.fill.solid()
                shp.fill.fore_color.rgb = ACCENT if k_rev == 1 else TEAL
                shp.line.fill.background()
                _, tf = textbox(s, x + Inches(0.15), Inches(3.3), cw - Inches(0.3), Inches(1.0),
                                anchor=MSO_ANCHOR.MIDDLE)
                para(tf, f"گام {k_rev + 1}", size=16, bold=True, color=CARD, first=True,
                     align=PP_ALIGN.CENTER)
            notes(s, "شه‌پرون RTL: گام اول راست‌ترین. فقط گام محوری یک رنگ متفاوت.")
        elif pat == "premise":
            bg(s)
            header(s, "نتیجه را اینجا بنویسید (از مقدمه راست نتیجه می‌شود)",
                   kicker="مقدمه ← نتیجه", num=num, total=total)
            _, tf = textbox(s, Inches(7.0), Inches(2.2), Inches(5.7), Inches(3.6))
            para(tf, "مقدمه (راست)", size=15, bold=True, color=TEAL, first=True, space_after=8)
            para(tf, "متن نمونه: جدول/واقعیت بالادستی.", size=14)
            _, tf2 = textbox(s, MARGIN, Inches(2.2), Inches(5.0), Inches(3.6))
            para(tf2, "نتیجه (چپ)", size=15, bold=True, color=TEAL, first=True, space_after=8)
            para(tf2, "متن نمونه: پس چه می‌شود — تیتر مستقل اسم‌عبارت.", size=14)
            tri = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                     Inches(6.3), Inches(3.6), Inches(0.5), Inches(0.5))
            tri.fill.solid()
            tri.fill.fore_color.rgb = ACCENT
            tri.line.fill.background()
            try:
                tri.rotation = 270.0
            except Exception:
                pass
            notes(s, "مقدمه→نتیجه با تک‌مثلث وسط (§4.27). تیتر چپ اسم‌عبارت مستقل.")
        elif pat == "decision":
            bg(s)
            header(s, "چه تصمیمی می‌خواهید؟ (اینجا بنویسید)", kicker="تصمیم", num=num, total=total)
            _, tf = textbox(s, MARGIN, Inches(1.9), SLIDE_W - 2 * MARGIN, Inches(4.4))
            para(tf, "متن نمونه: تصمیم + پیش‌فرض‌ها + خواسته‌ها — سه بلوک.", size=15, first=True)
            notes(s, "صفحه تصمیم: چه چیزی، بر چه مبنایی، چه کمکی لازم است.")
        elif pat == "stat":
            bg(s)
            header(s, "عدد بزرگ + خوانش آن", kicker="شاخص", num=num, total=total)
            _, tf = textbox(s, MARGIN, Inches(2.6), SLIDE_W - 2 * MARGIN, Inches(2.4),
                            anchor=MSO_ANCHOR.MIDDLE)
            para(tf, "۲۶", size=64, bold=True, color=ACCENT, first=True,
                 align=PP_ALIGN.CENTER, rtl=False)
            para(tf, "متن نمونه: بزرگی و معنای عدد در یک خط.", size=16,
                 align=PP_ALIGN.CENTER)
            notes(s, "عدد بزرگ: یک عدد + یک خوانش کوتاه؛ ارقام فارسی طبق سیاست دک.")
        else:  # map, toc, vsteps, matrix, gantt, case + future ids
            bg(s)
            labels = {"map": ("نقشه کلی (هر سطر → P.n)", "فقط دک‌های بالای ~۱۰ اسلاید"),
                      "toc": ("فهرست (تک‌ستون)", "یک ستون"),
                      "vsteps": ("گام‌های عمودی + نظر کناری", "راست گام‌ها، چپ نظر"),
                      "matrix": ("ماتریس ۲×۲", "خاکستری برای خنثی، هایلایت برای کانون"),
                      "gantt": ("گانت", "سطر = نقش، ستون = مقطع"),
                      "case": ("چالش (راست) / راه‌حل (چپ)", "راه‌حل با محور سطری")}
            head, kick = labels.get(pat, (pat, "جایگزین کنید"))
            header(s, head, kicker=kick, num=num, total=total)
            _, tf = textbox(s, MARGIN, Inches(1.9), SLIDE_W - 2 * MARGIN, Inches(4.4))
            for j in range(3):
                para(tf, f"متن نمونه {j + 1} — با خط داستانی جایگزین کنید",
                     size=15, first=(j == 0), bullet=True, space_after=8)
            notes(s, f"الگوی {pat}. نمونه‌ها را کامل جایگزین کنید (‎§2.8 — بدون 자리).")

    prs.save(out)
    print(f"wrote {out}: {len(parts)} slides (numbered {total})")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--parts", help="comma list, e.g. cover,map,axis,chart,decision,back")
    ap.add_argument("--title", default="نام سند")
    ap.add_argument("-o", "--out", default="mydeck.pptx")
    ap.add_argument("--list", action="store_true", help="list pattern IDs and exit")
    a = ap.parse_args()
    if a.list or not a.parts:
        if not a.parts and not a.list:
            print("--parts … (e.g. --parts cover,map,axis,chart,decision,back). Available:", file=sys.stderr)
        for pid, desc in PATTERNS:
            print(f"{pid:10s} {desc}")
        return
    parts = [t.strip().lower() for t in a.parts.split(",")]
    ids = {p for p, _ in PATTERNS}
    for t in parts:
        if t not in ids:
            sys.exit(f"unknown pattern: {t} (see --list)")
    build(parts, a.title, a.out)


if __name__ == "__main__":
    main()
