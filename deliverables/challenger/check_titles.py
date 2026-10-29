# -*- coding: utf-8 -*-
"""Title QA: re-wrap every title with HarfBuzz and flag orphan last lines.

Rule (slide-rules-fa §2.1/§2.13): max 2 lines, break at a meaning boundary,
never leave a 1-3 character orphan on the last line.
"""
import os
import sys

from pptx import Presentation
from pptx.util import Emu

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import shaping


def wrap(text, bold, size, max_w):
    words = [w for w in text.split(" ") if w]
    lines, cur = [], []
    for w in words:
        trial = " ".join(cur + [w])
        if cur and shaping.measure_in(trial, bold, size) > max_w:
            lines.append(" ".join(cur))
            cur = [w]
        else:
            cur.append(w)
    if cur:
        lines.append(" ".join(cur))
    return lines


def main(path):
    prs = Presentation(path)
    problems = 0
    for i, slide in enumerate(prs.slides, start=1):
        biggest = None
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                runs = [r for r in p.runs if r.text.strip()]
                if not runs:
                    continue
                size = runs[0].font.size.pt if runs[0].font.size else 18
                if size >= 26 and (biggest is None or size > biggest[0]):
                    biggest = (size, "".join(r.text for r in runs), Emu(sh.width).inches,
                               bool(runs[0].font.bold))
        if not biggest:
            continue
        size, text, w, bold = biggest
        lines = wrap(text, bold, size, w)
        last = lines[-1] if lines else ""
        flag = ""
        if len(lines) > 2:
            flag = "FAIL >2 lines"
        elif len(lines) == 2 and len(last.split()) <= 1:
            flag = "FAIL orphan word on line 2"
        elif len(lines) == 2 and len(last) <= 3:
            flag = "FAIL tiny last line"
        if flag:
            problems += 1
        print(f"S{i:02d} [{size:.0f}pt] lines={len(lines)} last='{last}' "
              f"{flag}  | {text[:52]}")
    print(f"\n{problems} title problem(s)")
    return problems


if __name__ == "__main__":
    p = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "Dodge-Challenger-FA.pptx")
    sys.exit(1 if main(p) else 0)
