# -*- coding: utf-8 -*-
"""Visual QA renderer for the .pptx — real HarfBuzz shaping, matplotlib output.

Not pixel-identical to PowerPoint, but faithful for text shaping, alignment,
line breaks, overlaps, clipping and whitespace balance. All units in inches.

Usage: python render_preview.py [slide numbers...]
"""
import os
import sys
from io import BytesIO

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Ellipse, PathPatch, Rectangle
from matplotlib.path import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import shaping

HERE = os.path.dirname(os.path.abspath(__file__))
DPI = 110
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def inch(emu):
    return Emu(emu).inches


def color_of(fill_el):
    srgb = fill_el.find(A + "srgbClr")
    if srgb is None:
        return None, 1.0
    v = srgb.get("val")
    rgb = tuple(int(v[i:i + 2], 16) / 255 for i in (0, 2, 4))
    a = srgb.find(A + "alpha")
    return rgb, (1.0 if a is None else int(a.get("val")) / 100000)


def solid_of(shape):
    spPr = shape._element.spPr
    if spPr is None:
        return None, 1.0
    sf = spPr.find(A + "solidFill")
    return (None, 1.0) if sf is None else color_of(sf)


def wrap_words(words, max_w):
    lines, cur, cur_w = [], [], 0.0
    for w in words:
        ww = shaping.measure_in(w[0], w[2], w[1])
        sp = shaping.measure_in(" ", w[2], w[1]) if cur else 0.0
        if cur and cur_w + sp + ww > max_w + 0.002:
            lines.append(cur)
            cur, cur_w, sp = [], 0.0, 0.0
        cur.append(w)
        cur_w += sp + ww
    if cur:
        lines.append(cur)
    return lines


def _is_ar(word):
    return any(shaping.is_arabic_letter(ch) for ch in word[0])


def _runs(words):
    """Group consecutive words into runs of one script class."""
    runs, cur = [], []
    for w in words:
        cls = "ar" if _is_ar(w) else "ltr"
        if cur and cur[0][0] != cls:
            runs.append((cur[0][0], [x[1] for x in cur]))
            cur = []
        cur.append((cls, w))
    if cur:
        runs.append((cur[0][0], [x[1] for x in cur]))
    return runs


def draw_words(ax, words, bx, bw, baseline, align):
    """Mini-bidi: Arabic runs read right-to-left, digits/Latin runs left-to-right."""
    if not words:
        return
    space = shaping.measure_in(" ", False, words[0][1])
    left_s = shaping.measure_in(" ", False, words[0][1])
    runs = _runs(words)
    run_widths = []
    for cls, ws in runs:
        wid = sum(shaping.measure_in(w[0], w[2], w[1]) for w in ws) + left_s * (len(ws) - 1)
        run_widths.append(wid)
    total = sum(run_widths) + space * (len(runs) - 1)

    if align == "center":
        block_left = bx + (bw - total) / 2
    elif align == "left":
        block_left = bx
    else:
        block_left = bx + bw - total

    para_rtl = any(cls == "ar" for cls, _ in runs)
    if para_rtl:
        cursor = block_left + total          # right edge, first logical run is rightmost
        order = list(range(len(runs)))
    else:
        cursor = block_left                  # pure LTR line
        order = list(range(len(runs)))

    for i in order:
        cls, ws = runs[i]
        rw = run_widths[i]
        if para_rtl:
            span_left = cursor - rw
        else:
            span_left = cursor
        x = span_left
        if cls == "ar":
            # right-to-left inside the run
            for w in ws:
                wd = shaping.measure_in(w[0], w[2], w[1])
                right = span_left + rw
                # place incrementally from the right
                x = right
                break
            cur = span_left + rw
            for w in ws:
                wd = shaping.measure_in(w[0], w[2], w[1])
                cur -= wd
                _emit(ax, w, cur, baseline)
                cur -= left_s
        else:
            cur = span_left
            for w in ws:
                _emit(ax, w, cur, baseline)
                cur += shaping.measure_in(w[0], w[2], w[1]) + left_s
        cursor = span_left - space if para_rtl else span_left + rw + space


def _emit(ax, w, x, baseline):
    for verts, codes in shaping.text_paths_in(w[0], w[2], w[1], x, baseline):
        c = (w[3][0] / 255, w[3][1] / 255, w[3][2] / 255) if w[3] else (0.96, 0.96, 0.97)
        ax.add_patch(PathPatch(Path(verts, codes), facecolor=c, edgecolor="none", zorder=3))


def render(pptx_path, outdir, only=None):
    prs = Presentation(pptx_path)
    W, H = inch(prs.slide_width), inch(prs.slide_height)
    os.makedirs(outdir, exist_ok=True)
    out_files = []
    for idx, slide in enumerate(prs.slides, start=1):
        if only and idx not in only:
            continue
        fig = plt.figure(figsize=(W, H), dpi=DPI)
        ax = fig.add_axes([0, 0, 1, 1])
        ax.set_xlim(0, W)
        ax.set_ylim(H, 0)
        ax.axis("off")
        bg = (10 / 255, 11 / 255, 13 / 255)
        bg_el = slide._element.find(".//" + A + "bg")
        if bg_el is not None:
            rgb, _ = color_of(bg_el)
            bg = rgb or bg
        ax.add_patch(Rectangle((0, 0), W, H, color=bg, zorder=0))

        for sh in slide.shapes:
            if sh.shape_type is not None and sh.shape_type == 13:
                pic = Image.open(BytesIO(sh.image.blob)).convert("RGB")
                iw, ih = pic.size
                box = (int((sh.crop_left or 0) * iw), int((sh.crop_top or 0) * ih),
                       int((1 - (sh.crop_right or 0)) * iw),
                       int((1 - (sh.crop_bottom or 0)) * ih))
                ax.imshow(pic.crop(box),
                          extent=(inch(sh.left), inch(sh.left) + inch(sh.width),
                                  inch(sh.top) + inch(sh.height), inch(sh.top)),
                          aspect="auto", zorder=1)
                continue

            rgb, alpha = solid_of(sh)
            if rgb:
                x, y = inch(sh.left), inch(sh.top)
                w, h = inch(sh.width), inch(sh.height)
                if sh.shape_type == 9 or "oval" in (sh.name or "").lower():
                    ax.add_patch(Ellipse((x + w / 2, y + h / 2), w, h, color=rgb,
                                         alpha=alpha, zorder=2))
                else:
                    ax.add_patch(Rectangle((x, y), w, h, color=rgb, alpha=alpha, zorder=2))

            if not sh.has_text_frame:
                continue
            tf = sh.text_frame
            bx, by = inch(sh.left), inch(sh.top)
            bw, bh = inch(sh.width), inch(sh.height)

            paras = []
            for p in tf.paragraphs:
                words = []
                for r in p.runs:
                    size = r.font.size.pt if r.font.size else 18
                    bold = bool(r.font.bold)
                    col = None
                    try:
                        if r.font.color and r.font.color.rgb:
                            col = tuple(int(str(r.font.color.rgb)[i:i + 2], 16)
                                        for i in (0, 2, 4))
                    except Exception:
                        pass
                    for w in r.text.split(" "):
                        if w:
                            words.append((w, size, bold, col))
                if not words:
                    paras.append(None)
                    continue
                max_size = max(w[1] for w in words)
                ls = p.line_spacing or 1.2
                align = str(p.alignment) if p.alignment is not None else "RIGHT"
                al = "center" if "CENTER" in align else ("left" if "LEFT" in align else "right")
                paras.append((wrap_words(words, bw), ls, max_size, al))

            def pheight(p):
                return len(p[0]) * (p[2] * p[1] / 72) if p else 0.18

            total = sum(pheight(p) for p in paras)
            anchor = str(tf.vertical_anchor or "TOP")
            y = by + ((bh - total) / 2 if "MIDDLE" in anchor else 0)

            for p in paras:
                if not p:
                    y += 0.18
                    continue
                lines, ls, max_size, al = p
                lh = max_size * ls / 72
                for words in lines:
                    baseline = y + lh * 0.74
                    draw_words(ax, words, bx, bw, baseline, al)
                    y += lh

        out = os.path.join(outdir, f"slide-{idx:02d}.png")
        fig.savefig(out, dpi=DPI)
        plt.close(fig)
        out_files.append(out)
    return out_files


if __name__ == "__main__":
    nums = [int(a) for a in sys.argv[1:]] or None
    print("\n".join(render(os.path.join(HERE, "Dodge-Challenger-FA.pptx"),
                          os.path.join(HERE, "preview"), nums)))
